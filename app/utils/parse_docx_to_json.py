from docx import Document
from docx.oxml.ns import qn
import json


def extract_with_structure(docx_path):
    doc = Document(docx_path)
    structured_content = []

    for para in doc.paragraphs:
        # 스타일 정보 포함
        content = {
            'text': para.text,
            'style': para.style.name,  # Heading 1, Normal 등
            'level': para.style.name if 'Heading' in para.style.name else None
        }
        structured_content.append(content)

    # 테이블 처리
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            table_data.append(row_data)
        structured_content.append({'type': 'table', 'data': table_data})

    return structured_content