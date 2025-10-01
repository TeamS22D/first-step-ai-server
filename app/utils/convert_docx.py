from docx import Document


def convert_docx(file_path: str):
    document = Document(file_path)
    text_content = []
    for para in document.paragraphs:
        text_content.append(para.text)

    return '\n'.join(text_content)


def convert_docx_table(file_path: str):
    document = Document(file_path)
    table_content = []

    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)

        table_content.append(rows)

    return table_content

def parse_styles(file_path: str):
    document = Document(file_path)

    styles_content = []

    for style in document.styles:

        styles_content.append(style.name)

    for para in document.paragraphs:
        styles_content.append(para.style)


    return styles_content